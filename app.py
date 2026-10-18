import os
import json
import uuid
import shutil
from typing import List, Optional, Dict, Any
from datetime import datetime
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Header, Depends
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

from auth import (
    hash_pw, get_admin_config, save_admin_config,
    get_users, save_users, get_subjects, save_subjects,
    create_session, get_session, delete_session
)
from parsers import parse_lecture_document, parse_exam_questions
from matcher import LectureExamMatcher

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
LECTURES_UPLOAD_DIR = os.path.join(UPLOADS_DIR, "lectures")
EXAMS_UPLOAD_DIR = os.path.join(UPLOADS_DIR, "exams")

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(LECTURES_UPLOAD_DIR, exist_ok=True)
os.makedirs(EXAMS_UPLOAD_DIR, exist_ok=True)

LECTURES_DB_FILE = os.path.join(DATA_DIR, "lectures.json")
EXAMS_DB_FILE = os.path.join(DATA_DIR, "exams.json")
RESULTS_DB_FILE = os.path.join(DATA_DIR, "results.json")

def read_db(filepath):
    if os.path.exists(filepath):
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return []
    return []

def write_db(filepath, data):
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

app = FastAPI(title="족보-강의자료 선지 매칭 포털 (과목별 & 권한 관리)")
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")

# ----------------- 인증 의존성 -----------------

def get_current_session(authorization: Optional[str] = Header(None)) -> Dict[str, Any]:
    if not authorization:
        raise HTTPException(status_code=401, detail="로그인이 필요합니다.")
    token = authorization.replace("Bearer ", "").strip()
    session = get_session(token)
    if not session:
        raise HTTPException(status_code=401, detail="세션이 만료되었거나 유효하지 않습니다. 다시 로그인해주세요.")
    return session

def require_admin(session: Dict[str, Any] = Depends(get_current_session)) -> Dict[str, Any]:
    if session.get("role") != "admin":
        raise HTTPException(status_code=403, detail="관리자 권한이 필요합니다.")
    return session

def require_approved_user(session: Dict[str, Any] = Depends(get_current_session)) -> Dict[str, Any]:
    if session.get("role") == "admin":
        return session
    users = get_users()
    user = next((u for u in users if u["id"] == session.get("user_id")), None)
    if not user or user.get("status") != "approved":
        raise HTTPException(status_code=403, detail="관리자의 승인이 필요한 계정입니다.")
    return session

# ----------------- 로그인 / 회원가입 API -----------------

class RegisterRequest(BaseModel):
    username: str
    password: str
    name: str
    note: Optional[str] = ""

@app.post("/api/auth/register")
def register_user(req: RegisterRequest):
    users = get_users()
    if any(u["username"] == req.username for u in users):
        raise HTTPException(status_code=400, detail="이미 등록된 아이디입니다.")
    
    new_user = {
        "id": "u_" + str(uuid.uuid4())[:8],
        "username": req.username.strip(),
        "password_hash": hash_pw(req.password),
        "name": req.name.strip(),
        "status": "pending", # 기본 승인 대기
        "role": "user",
        "note": req.note.strip() if req.note else "",
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M")
    }
    users.append(new_user)
    save_users(users)
    return {
        "success": True, 
        "message": "가입 신청이 완료되었습니다! 관리자가 승인한 후 이용하실 수 있습니다."
    }

class LoginRequest(BaseModel):
    username: str
    password: str

@app.post("/api/auth/login")
def user_login(req: LoginRequest):
    users = get_users()
    target = next((u for u in users if u["username"] == req.username.strip()), None)
    
    if not target or target.get("password_hash") != hash_pw(req.password):
        raise HTTPException(status_code=401, detail="아이디 또는 비밀번호가 일치하지 않습니다.")
        
    if target.get("status") == "pending":
        raise HTTPException(status_code=403, detail="관리자의 가입 승인 대기 중입니다. 승인 후 로그인해주세요.")
    elif target.get("status") == "rejected":
        raise HTTPException(status_code=403, detail="이용 승인이 거절된 계정입니다. 관리자에게 문의하세요.")
        
    token = create_session({
        "user_id": target["id"],
        "username": target["username"],
        "name": target["name"],
        "role": "user"
    })
    return {
        "success": True,
        "token": token,
        "user": {
            "id": target["id"],
            "username": target["username"],
            "name": target["name"],
            "role": "user"
        }
    }

@app.post("/api/auth/admin-login")
def admin_login(req: LoginRequest):
    admin_cfg = get_admin_config()
    if req.username.strip() != admin_cfg.get("username") or hash_pw(req.password) != admin_cfg.get("password_hash"):
        raise HTTPException(status_code=401, detail="관리자 아이디 또는 비밀번호가 일치하지 않습니다.")
        
    token = create_session({
        "user_id": "admin",
        "username": admin_cfg["username"],
        "name": "시스템 관리자",
        "role": "admin"
    })
    return {
        "success": True,
        "token": token,
        "user": {
            "id": "admin",
            "username": admin_cfg["username"],
            "name": "시스템 관리자",
            "role": "admin"
        }
    }

@app.get("/api/auth/me")
def get_me(session: Dict[str, Any] = Depends(get_current_session)):
    return {"authenticated": True, "user": session}

@app.post("/api/auth/logout")
def logout(authorization: Optional[str] = Header(None)):
    if authorization:
        token = authorization.replace("Bearer ", "").strip()
        delete_session(token)
    return {"success": True}

class ChangePasswordRequest(BaseModel):
    old_password: str
    new_password: str

@app.post("/api/auth/change-password")
def change_user_password(req: ChangePasswordRequest, session: Dict[str, Any] = Depends(require_approved_user)):
    users = get_users()
    user_id = session.get("user_id")
    target = next((u for u in users if u["id"] == user_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="사용자 정보를 찾을 수 없습니다.")
    if target.get("password_hash") != hash_pw(req.old_password):
        raise HTTPException(status_code=400, detail="현재 비밀번호가 일치하지 않습니다.")
    if len(req.new_password.strip()) < 4:
        raise HTTPException(status_code=400, detail="새 비밀번호는 최소 4자 이상이어야 합니다.")
    target["password_hash"] = hash_pw(req.new_password.strip())
    save_users(users)
    return {"success": True, "message": "비밀번호가 성공적으로 변경되었습니다."}

class AdminResetPasswordRequest(BaseModel):
    new_password: str

@app.post("/api/admin/users/{user_id}/reset-password")
def admin_reset_user_password(user_id: str, req: AdminResetPasswordRequest, admin: Dict[str, Any] = Depends(require_admin)):
    users = get_users()
    target = next((u for u in users if u["id"] == user_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="사용자를 찾을 수 없습니다.")
    target["password_hash"] = hash_pw(req.new_password.strip())
    save_users(users)
    return {"success": True, "message": f"{target.get('name')} 사용자의 비밀번호가 재설정되었습니다."}


# ----------------- 관리자: 이용자 승인/관리 API -----------------

@app.get("/api/admin/users")
def list_admin_users(admin: Dict[str, Any] = Depends(require_admin)):
    users = get_users()
    # 비밀번호 해시는 제외하고 응답
    safe_users = []
    for u in users:
        item = {k: v for k, v in u.items() if k != "password_hash"}
        safe_users.append(item)
    return safe_users

class UserStatusUpdate(BaseModel):
    status: str # approved | rejected | pending

@app.post("/api/admin/users/{user_id}/status")
def update_user_status(user_id: str, req: UserStatusUpdate, admin: Dict[str, Any] = Depends(require_admin)):
    users = get_users()
    target = next((u for u in users if u["id"] == user_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="사용자를 찾을 수 없습니다.")
    target["status"] = req.status
    save_users(users)
    return {"success": True, "user_id": user_id, "status": req.status}

class AdminCreateUserRequest(BaseModel):
    username: str
    password: str
    name: str
    note: Optional[str] = ""

@app.post("/api/admin/users/create")
def admin_create_user(req: AdminCreateUserRequest, admin: Dict[str, Any] = Depends(require_admin)):
    users = get_users()
    if any(u["username"] == req.username for u in users):
        raise HTTPException(status_code=400, detail="이미 존재하는 아이디입니다.")
    new_user = {
        "id": "u_" + str(uuid.uuid4())[:8],
        "username": req.username.strip(),
        "password_hash": hash_pw(req.password),
        "name": req.name.strip(),
        "status": "approved", # 관리자가 직접 추가 시 즉시 승인
        "role": "user",
        "note": req.note or "관리자 직접 생성",
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M")
    }
    users.append(new_user)
    save_users(users)
    return {"success": True, "user": {k: v for k, v in new_user.items() if k != "password_hash"}}

@app.delete("/api/admin/users/{user_id}")
def delete_user(user_id: str, admin: Dict[str, Any] = Depends(require_admin)):
    users = get_users()
    users = [u for u in users if u["id"] != user_id]
    save_users(users)
    return {"success": True}

class AdminPasswordChange(BaseModel):
    old_password: str
    new_password: str

@app.post("/api/admin/change-password")
def change_admin_password(req: AdminPasswordChange, admin: Dict[str, Any] = Depends(require_admin)):
    cfg = get_admin_config()
    if hash_pw(req.old_password) != cfg.get("password_hash"):
        raise HTTPException(status_code=400, detail="기존 관리자 비밀번호가 일치하지 않습니다.")
    cfg["password_hash"] = hash_pw(req.new_password)
    save_admin_config(cfg)
    return {"success": True, "message": "관리자 비밀번호가 변경되었습니다."}

@app.get("/api/admin/stats")
def get_admin_stats(admin: Dict[str, Any] = Depends(require_admin)):
    users = get_users()
    subs = get_subjects()
    lectures = read_db(LECTURES_DB_FILE)
    exams = read_db(EXAMS_DB_FILE)
    results = read_db(RESULTS_DB_FILE)
    
    return {
        "total_users": len(users),
        "pending_users": len([u for u in users if u.get("status") == "pending"]),
        "approved_users": len([u for u in users if u.get("status") == "approved"]),
        "total_subjects": len(subs),
        "total_lectures": len(lectures),
        "total_exams": len(exams),
        "total_results": len(results)
    }

# ----------------- 과목(Subject) 관리 API -----------------

@app.get("/api/subjects")
def list_subjects(user: Dict[str, Any] = Depends(require_approved_user)):
    subs = get_subjects()
    lectures = read_db(LECTURES_DB_FILE)
    exams = read_db(EXAMS_DB_FILE)
    
    # 각 과목별 등록된 자료 수 집계
    enhanced = []
    for s in subs:
        s_id = s["id"]
        l_cnt = len([l for l in lectures if l.get("subject_id") == s_id])
        e_cnt = len([e for e in exams if e.get("subject_id") == s_id])
        enhanced.append({
            **s,
            "lecture_count": l_cnt,
            "exam_count": e_cnt
        })
    return enhanced

class SubjectCreateRequest(BaseModel):
    name: str
    description: Optional[str] = ""

@app.post("/api/admin/subjects")
def create_subject(req: SubjectCreateRequest, admin: Dict[str, Any] = Depends(require_admin)):
    subs = get_subjects()
    new_sub = {
        "id": "sub_" + str(uuid.uuid4())[:8],
        "name": req.name.strip(),
        "description": req.description.strip() if req.description else "",
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M")
    }
    subs.append(new_sub)
    save_subjects(subs)
    return {"success": True, "subject": new_sub}

@app.put("/api/admin/subjects/{subject_id}")
def update_subject(subject_id: str, req: SubjectCreateRequest, admin: Dict[str, Any] = Depends(require_admin)):
    subs = get_subjects()
    target = next((s for s in subs if s["id"] == subject_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="과목을 찾을 수 없습니다.")
    target["name"] = req.name.strip()
    target["description"] = req.description.strip()
    save_subjects(subs)
    return {"success": True, "subject": target}

@app.delete("/api/admin/subjects/{subject_id}")
def delete_subject(subject_id: str, admin: Dict[str, Any] = Depends(require_admin)):
    subs = get_subjects()
    subs = [s for s in subs if s["id"] != subject_id]
    save_subjects(subs)
    return {"success": True}

# ----------------- 강의 자료 관리 API (과목별) -----------------

@app.get("/api/lectures")
def list_lectures(subject_id: Optional[str] = None, user: Dict[str, Any] = Depends(require_approved_user)):
    docs = read_db(LECTURES_DB_FILE)
    if subject_id:
        docs = [d for d in docs if d.get("subject_id") == subject_id]
        
    summary = []
    for d in docs:
        summary.append({
            "id": d["id"],
            "subject_id": d.get("subject_id"),
            "filename": d["filename"],
            "unit_count": len(d.get("units", [])),
            "unit_type": d.get("unit_type", "슬라이드/페이지"),
            "uploaded_at": d.get("uploaded_at"),
            "uploaded_by": d.get("uploaded_by", "사용자")
        })
    return summary

@app.post("/api/lectures/upload")
async def upload_lectures(
    subject_id: str = Form(...),
    files: List[UploadFile] = File(...),
    user: Dict[str, Any] = Depends(require_approved_user)
):
    docs = read_db(LECTURES_DB_FILE)
    uploaded_items = []
    
    for file in files:
        file_id = str(uuid.uuid4())[:8]
        safe_filename = file.filename
        save_path = os.path.join(LECTURES_UPLOAD_DIR, f"{file_id}_{safe_filename}")
        
        with open(save_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
            
        units = parse_lecture_document(save_path)
        unit_type = units[0]["unit_type"] if units else "단위"
        
        for u in units:
            u["doc_id"] = file_id
            u["doc_title"] = safe_filename
            
        doc_entry = {
            "id": file_id,
            "subject_id": subject_id,
            "filename": safe_filename,
            "path": save_path,
            "unit_type": unit_type,
            "units": units,
            "uploaded_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "uploaded_by": user.get("name", "사용자")
        }
        docs.append(doc_entry)
        uploaded_items.append({
            "id": file_id,
            "filename": safe_filename,
            "unit_count": len(units),
            "unit_type": unit_type
        })
        
    write_db(LECTURES_DB_FILE, docs)
    return {"success": True, "uploaded": uploaded_items}

@app.delete("/api/lectures/{doc_id}")
def delete_lecture(doc_id: str, user: Dict[str, Any] = Depends(require_approved_user)):
    docs = read_db(LECTURES_DB_FILE)
    target = next((d for d in docs if d["id"] == doc_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="강의 자료를 찾을 수 없습니다.")
    
    if os.path.exists(target.get("path", "")):
        try:
            os.remove(target["path"])
        except Exception:
            pass
            
    docs = [d for d in docs if d["id"] != doc_id]
    write_db(LECTURES_DB_FILE, docs)
    return {"success": True, "message": "강의 자료가 삭제되었습니다."}

@app.get("/api/lectures/{doc_id}/units")
def get_lecture_units(doc_id: str, user: Dict[str, Any] = Depends(require_approved_user)):
    docs = read_db(LECTURES_DB_FILE)
    target = next((d for d in docs if d["id"] == doc_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="자료를 찾을 수 없습니다.")
    return {
        "id": target["id"],
        "filename": target["filename"],
        "unit_type": target["unit_type"],
        "units": target.get("units", [])
    }

# ----------------- 족보 문항 관리 API (과목별) -----------------

@app.get("/api/exams")
def list_exams(subject_id: Optional[str] = None, user: Dict[str, Any] = Depends(require_approved_user)):
    exams = read_db(EXAMS_DB_FILE)
    if subject_id:
        exams = [e for e in exams if e.get("subject_id") == subject_id]
    return [{
        "id": e["id"],
        "subject_id": e.get("subject_id"),
        "title": e["title"],
        "question_count": len(e.get("questions", [])),
        "created_at": e.get("created_at"),
        "created_by": e.get("created_by", "사용자")
    } for e in exams]

class ExamPasteRequest(BaseModel):
    subject_id: str
    title: str
    text_content: str

@app.post("/api/exams/paste")
def paste_exam(req: ExamPasteRequest, user: Dict[str, Any] = Depends(require_approved_user)):
    exam_id = str(uuid.uuid4())[:8]
    questions = parse_exam_questions(req.text_content)
    
    exam_entry = {
        "id": exam_id,
        "subject_id": req.subject_id,
        "title": req.title or "붙여넣은 족보 문항",
        "raw_text": req.text_content,
        "questions": questions,
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "created_by": user.get("name", "사용자")
    }
    
    exams = read_db(EXAMS_DB_FILE)
    exams.append(exam_entry)
    write_db(EXAMS_DB_FILE, exams)
    
    return {"success": True, "id": exam_id, "question_count": len(questions)}

@app.post("/api/exams/upload")
async def upload_exam(
    subject_id: str = Form(...),
    title: Optional[str] = Form(None),
    file: UploadFile = File(...),
    user: Dict[str, Any] = Depends(require_approved_user)
):
    exam_id = str(uuid.uuid4())[:8]
    safe_name = file.filename
    save_path = os.path.join(EXAMS_UPLOAD_DIR, f"{exam_id}_{safe_name}")
    
    with open(save_path, "wb") as buf:
        shutil.copyfileobj(file.file, buf)
        
    units = parse_lecture_document(save_path)
    full_text = "\n\n".join([u["text"] for u in units])
    questions = parse_exam_questions(full_text)
    
    exam_entry = {
        "id": exam_id,
        "subject_id": subject_id,
        "title": title or safe_name,
        "filename": safe_name,
        "questions": questions,
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "created_by": user.get("name", "사용자")
    }
    
    exams = read_db(EXAMS_DB_FILE)
    exams.append(exam_entry)
    write_db(EXAMS_DB_FILE, exams)
    
    return {"success": True, "id": exam_id, "question_count": len(questions)}

@app.get("/api/exams/{exam_id}")
def get_exam(exam_id: str, user: Dict[str, Any] = Depends(require_approved_user)):
    exams = read_db(EXAMS_DB_FILE)
    target = next((e for e in exams if e["id"] == exam_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="족보 세트를 찾을 수 없습니다.")
    return target

@app.delete("/api/exams/{exam_id}")
def delete_exam(exam_id: str, user: Dict[str, Any] = Depends(require_approved_user)):
    exams = read_db(EXAMS_DB_FILE)
    exams = [e for e in exams if e["id"] != exam_id]
    write_db(EXAMS_DB_FILE, exams)
    return {"success": True}

# ----------------- 문제은행 기출 문제집 & 풀이 모드 API -----------------

HISTORY_DB_FILE = os.path.join(DATA_DIR, "solving_history.json")

class SaveAnswerKeysRequest(BaseModel):
    answer_keys: Dict[str, str]

@app.post("/api/workbooks/{exam_id}/answers")
def save_answer_keys(exam_id: str, req: SaveAnswerKeysRequest, user: Dict[str, Any] = Depends(require_approved_user)):
    exams = read_db(EXAMS_DB_FILE)
    target = next((e for e in exams if e["id"] == exam_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="문제집을 찾을 수 없습니다.")
    target["answer_keys"] = req.answer_keys
    write_db(EXAMS_DB_FILE, exams)
    return {"success": True, "message": "정답이 저장되었습니다."}

CIRCLED_TO_NUM = {
    '①': '1', '②': '2', '③': '3', '④': '4', '⑤': '5',
    '⑥': '6', '⑦': '7', '⑧': '8', '⑨': '9', '⑩': '10',
    '⑴': '1', '⑵': '2', '⑶': '3', '⑷': '4', '⑸': '5'
}

def normalize_ans_tokens(s: str) -> List[str]:
    if not s:
        return []
    res = str(s)
    for c, n in CIRCLED_TO_NUM.items():
        res = res.replace(c, n)
    return [t.strip().lower() for t in re.split(r'[,/·\s]+', res) if t.strip()]

def evaluate_exam_answer(user_ans: str, correct_key: str) -> Tuple[bool, bool]:
    """(has_key, is_correct) 반환"""
    if not correct_key or not correct_key.strip():
        return False, False
    u_trim = (user_ans or "").strip()
    c_trim = correct_key.strip()
    if not u_trim:
        return True, False
    
    # 1. 단순 텍스트 일치 (공백/대소문자 무시)
    if u_trim.lower() == c_trim.lower():
        return True, True
        
    # 2. 복수 정답 및 번호 토큰 비교 (예: "3,4", "1, 4", "③, ④")
    u_tokens = normalize_ans_tokens(u_trim)
    c_tokens = normalize_ans_tokens(c_trim)
    
    if len(u_tokens) == 1 and u_tokens[0] in c_tokens:
        return True, True
    if len(u_tokens) > 1 and sorted(u_tokens) == sorted(c_tokens):
        return True, True
        
    return True, False

class SubmitExamRequest(BaseModel):
    user_answers: Dict[str, str]
    time_spent_seconds: int

@app.post("/api/workbooks/{exam_id}/submit")
def submit_exam_attempt(exam_id: str, req: SubmitExamRequest, user: Dict[str, Any] = Depends(require_approved_user)):
    exams = read_db(EXAMS_DB_FILE)
    exam = next((e for e in exams if e["id"] == exam_id), None)
    if not exam:
        raise HTTPException(status_code=404, detail="문제집을 찾을 수 없습니다.")

    answer_keys = exam.get("answer_keys", {})
    questions = exam.get("questions", [])
    
    # 강의자료 매칭 정보 연계
    results = read_db(RESULTS_DB_FILE)
    matching_result = next((r for r in results if r.get("exam_id") == exam_id), None)
    
    if not matching_result:
        lectures = read_db(LECTURES_DB_FILE)
        subject_lectures = [l for l in lectures if l.get("subject_id") == exam.get("subject_id")]
        all_units = []
        for l in subject_lectures:
            for u in l.get("units", []):
                all_units.append({
                    "doc_id": l["id"],
                    "doc_title": l["filename"],
                    "unit_type": u.get("unit_type", "슬라이드"),
                    "unit_number": u.get("unit_number", 1),
                    "text": u.get("text", "")
                })
        if all_units:
            matcher = LectureExamMatcher(all_units)
            analyzed_questions = matcher.match_all_questions(questions)
        else:
            analyzed_questions = []
    else:
        analyzed_questions = matching_result.get("questions", [])

    total_q = len(questions)
    correct_count = 0
    scored_q = 0
    question_reviews = []

    for idx, q in enumerate(questions):
        q_num = str(q.get("q_number", idx + 1))
        user_choice = req.user_answers.get(q_num, "")
        correct_choice = answer_keys.get(q_num, "")
        
        has_key, is_correct = evaluate_exam_answer(user_choice, correct_choice)
        if has_key:
            scored_q += 1
            if is_correct:
                correct_count += 1
                
        matched_q = next((mq for mq in analyzed_questions if str(mq.get("q_number")) == q_num), None)
        
        question_reviews.append({
            "q_number": int(q_num) if q_num.isdigit() else idx + 1,
            "stem": q.get("stem", ""),
            "choices": q.get("choices", []),
            "user_choice": user_choice,
            "correct_choice": correct_choice,
            "is_correct": is_correct,
            "has_answer_key": has_key,
            "matched_analysis": matched_q.get("choices", []) if matched_q else []
        })

    score = round((correct_count / scored_q) * 100, 1) if scored_q > 0 else 0
    
    history = read_db(HISTORY_DB_FILE)
    attempt_id = str(uuid.uuid4())[:8]
    attempt_entry = {
        "id": attempt_id,
        "exam_id": exam_id,
        "exam_title": exam.get("title"),
        "subject_id": exam.get("subject_id"),
        "user_id": user.get("user_id"),
        "user_name": user.get("name"),
        "score": score,
        "correct_count": correct_count,
        "total_scored": scored_q,
        "total_questions": total_q,
        "time_spent_seconds": req.time_spent_seconds,
        "submitted_at": datetime.now().strftime("%Y-%m-%d %H:%M")
    }
    history.insert(0, attempt_entry)
    write_db(HISTORY_DB_FILE, history)

    return {
        "success": True,
        "attempt_id": attempt_id,
        "score": score,
        "correct_count": correct_count,
        "total_scored": scored_q,
        "total_questions": total_q,
        "time_spent_seconds": req.time_spent_seconds,
        "reviews": question_reviews
    }

@app.get("/api/workbooks/history")
def get_user_history(subject_id: Optional[str] = None, user: Dict[str, Any] = Depends(require_approved_user)):
    history = read_db(HISTORY_DB_FILE)
    user_id = user.get("user_id")
    if user.get("role") != "admin":
        history = [h for h in history if h.get("user_id") == user_id]
    if subject_id:
        history = [h for h in history if h.get("subject_id") == subject_id]
    return history


# ----------------- 선지 매칭 분석 API (과목별 격리) -----------------

class RunMatchRequest(BaseModel):
    subject_id: str
    exam_id: str
    lecture_ids: Optional[List[str]] = None

@app.post("/api/match/run")
def run_match(req: RunMatchRequest, user: Dict[str, Any] = Depends(require_approved_user)):
    exams = read_db(EXAMS_DB_FILE)
    exam = next((e for e in exams if e["id"] == req.exam_id), None)
    if not exam:
        raise HTTPException(status_code=404, detail="족보 세트를 찾을 수 없습니다.")
        
    lectures = read_db(LECTURES_DB_FILE)
    # 해당 과목의 강의자료만 필터링
    lectures = [l for l in lectures if l.get("subject_id") == req.subject_id]
    
    if req.lecture_ids:
        lectures = [l for l in lectures if l["id"] in req.lecture_ids]
        
    if not lectures:
        raise HTTPException(status_code=400, detail="이 과목에 등록된 강의 자료가 없습니다. 먼저 강의 자료를 등록해주세요.")

    all_units = []
    for l in lectures:
        for u in l.get("units", []):
            all_units.append({
                "doc_id": l["id"],
                "doc_title": l["filename"],
                "unit_type": u.get("unit_type", "슬라이드"),
                "unit_number": u.get("unit_number", 1),
                "text": u.get("text", "")
            })

    matcher = LectureExamMatcher(all_units)
    analyzed_questions = matcher.match_all_questions(exam.get("questions", []))
    
    total_choices = 0
    high_count = 0
    medium_count = 0
    low_count = 0
    for q in analyzed_questions:
        for c in q.get("choices", []):
            total_choices += 1
            st = c.get("analysis", {}).get("status", "low")
            if st == "high":
                high_count += 1
            elif st == "medium":
                medium_count += 1
            else:
                low_count += 1

    result_id = str(uuid.uuid4())[:8]
    result_entry = {
        "id": result_id,
        "subject_id": req.subject_id,
        "exam_id": req.exam_id,
        "exam_title": exam.get("title"),
        "lecture_count": len(lectures),
        "analyzed_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "analyzed_by": user.get("name", "사용자"),
        "summary": {
            "total_questions": len(analyzed_questions),
            "total_choices": total_choices,
            "high_count": high_count,
            "medium_count": medium_count,
            "low_count": low_count
        },
        "questions": analyzed_questions
    }
    
    results = read_db(RESULTS_DB_FILE)
    results.insert(0, result_entry)
    write_db(RESULTS_DB_FILE, results)
    
    return {"success": True, "result_id": result_id, "summary": result_entry["summary"]}

@app.get("/api/match/results")
def list_results(subject_id: Optional[str] = None, user: Dict[str, Any] = Depends(require_approved_user)):
    results = read_db(RESULTS_DB_FILE)
    if subject_id:
        results = [r for r in results if r.get("subject_id") == subject_id]
    return [{
        "id": r["id"],
        "subject_id": r.get("subject_id"),
        "exam_title": r.get("exam_title"),
        "analyzed_at": r.get("analyzed_at"),
        "summary": r.get("summary")
    } for r in results]

@app.get("/api/match/results/{result_id}")
def get_result(result_id: str, user: Dict[str, Any] = Depends(require_approved_user)):
    results = read_db(RESULTS_DB_FILE)
    target = next((r for r in results if r["id"] == result_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="분석 결과를 찾을 수 없습니다.")
    return target

# ----------------- 엑셀 다운로드 API -----------------

@app.get("/api/match/export/{result_id}/excel")
def export_excel(result_id: str):
    results = read_db(RESULTS_DB_FILE)
    target = next((r for r in results if r["id"] == result_id), None)
    if not target:
        raise HTTPException(status_code=404, detail="결과를 찾을 수 없습니다.")
        
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "선지 매칭 결과"
    
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="맑은 고딕", size=10)
    bold_font = Font(name="맑은 고딕", size=10, bold=True)
    
    fill_high = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_med = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    fill_low = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    headers = [
        "문제", "발문 (문제 내용)", "선지 번호", "선지 내용",
        "매칭 상태", "강의 자료명", "언급 위치 (슬라이드/페이지)",
        "일치도", "언급 문맥 (하이라이트 발췌)", "일치 키워드"
    ]
    
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    current_row = 2
    for q in target.get("questions", []):
        q_num = q.get("q_number")
        stem = q.get("stem", "")
        choices = q.get("choices", [])
        
        for c in choices:
            label = c.get("label", "")
            c_text = c.get("text", "")
            an = c.get("analysis", {})
            top = an.get("top_match") or {}
            
            status_badge = top.get("badge", "직접 언급 없음")
            doc_title = top.get("doc_title", "-")
            unit_loc = f"{top.get('unit_type', '위치')} {top.get('unit_number', '')}" if top else "-"
            score_str = f"{top.get('score', 0)}%" if top else "0%"
            snippet = top.get("snippet", "-") if top else "-"
            kws = ", ".join(top.get("matched_keywords", [])) if top else "-"
            
            row_data = [
                f"{q_num}번", stem, label, c_text,
                status_badge, doc_title, unit_loc,
                score_str, snippet, kws
            ]
            ws.append(row_data)
            
            status_val = top.get("status", "low")
            status_fill = fill_high if status_val == "high" else (fill_med if status_val == "medium" else fill_low)
            
            for col_i in range(1, len(headers) + 1):
                c_cell = ws.cell(row=current_row, column=col_i)
                c_cell.font = data_font
                c_cell.border = thin_border
                c_cell.alignment = Alignment(vertical="center", wrap_text=True)
                if col_i in [1, 3, 5, 7, 8]:
                    c_cell.alignment = Alignment(horizontal="center", vertical="center")
                if col_i == 5:
                    c_cell.fill = status_fill
                    c_cell.font = bold_font
            
            ws.row_dimensions[current_row].height = 36
            current_row += 1

    col_widths = {1: 10, 2: 35, 3: 10, 4: 35, 5: 16, 6: 25, 7: 18, 8: 12, 9: 45, 10: 22}
    for col_i, width in col_widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_i)].width = width

    export_path = os.path.join(DATA_DIR, f"matching_result_{result_id}.xlsx")
    wb.save(export_path)
    
    clean_title = target.get("exam_title", "matching_result").replace(" ", "_")
    return FileResponse(
        export_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"선지매칭결과_{clean_title}.xlsx"
    )

# ----------------- 페이지 서빙 -----------------

@app.get("/")
def serve_user_portal():
    return FileResponse(os.path.join(BASE_DIR, "static", "index.html"))

@app.get("/knumed2026jokbo")
@app.get("/knumed2026jokbo/")
def serve_knumed_user_portal():
    return FileResponse(os.path.join(BASE_DIR, "static", "index.html"))

@app.get("/admin")
def serve_admin_panel():
    return FileResponse(os.path.join(BASE_DIR, "static", "admin.html"))

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=True)
